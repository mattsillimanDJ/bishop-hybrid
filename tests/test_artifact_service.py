from docx import Document
from openpyxl import load_workbook

from app.services import artifact_service


def test_detects_docx_export_intent():
    request = artifact_service.detect_artifact_export_request(
        "make this a word doc: Launch Plan\n- write the intro\n- confirm budget"
    )

    assert request is not None
    assert request.kind == "docx"
    assert request.content == "Launch Plan\n- write the intro\n- confirm budget"


def test_detects_xlsx_export_intent():
    request = artifact_service.detect_artifact_export_request(
        "export this as xlsx: Name,Status\nBishop,Ready\nStemLab,Planning"
    )

    assert request is not None
    assert request.kind == "xlsx"
    assert request.content == "Name,Status\nBishop,Ready\nStemLab,Planning"


def test_creates_docx_file_from_simple_content(tmp_path):
    result = artifact_service.create_docx_artifact(
        content="Launch Plan\n\nSummary:\n- Draft intro\n- Confirm budget",
        output_dir=tmp_path,
    )

    assert result.kind == "docx"
    assert result.path.exists()
    assert result.filename.endswith(".docx")

    document = Document(result.path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    assert "Launch Plan" in paragraphs
    assert "Summary" in paragraphs
    assert "Draft intro" in paragraphs
    assert "Confirm budget" in paragraphs


def test_creates_xlsx_file_from_simple_table_content(tmp_path):
    result = artifact_service.create_xlsx_artifact(
        content="Name,Status\nBishop,Ready\nStemLab,Planning",
        output_dir=tmp_path,
    )

    assert result.kind == "xlsx"
    assert result.path.exists()
    assert result.filename.endswith(".xlsx")

    workbook = load_workbook(result.path)
    worksheet = workbook.active
    assert worksheet["A1"].value == "Name"
    assert worksheet["B1"].value == "Status"
    assert worksheet["A2"].value == "Bishop"
    assert worksheet["B2"].value == "Ready"
    assert worksheet["A1"].font.bold is True
    assert worksheet.column_dimensions["A"].width >= 12
